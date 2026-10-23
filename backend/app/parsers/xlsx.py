from datetime import datetime
from io import BytesIO
from typing import BinaryIO

from openpyxl import load_workbook

from app.parsers.base import DocumentParser
from app.parsers.models import ExtractedDocument, ExtractedElement


class XLSXParser(DocumentParser):
    """Parser for Microsoft Excel workbooks."""

    def parse(self, file: BinaryIO) -> ExtractedDocument:
        raw = file.read()

        # Formula view (what is typed in the cell) and value view (the result
        # Excel/LibreOffice last calculated and saved). Using the saved result
        # makes totals, lookups and flags searchable as real numbers.
        workbook = load_workbook(BytesIO(raw), read_only=False, data_only=False)
        values_workbook = load_workbook(BytesIO(raw), read_only=False, data_only=True)

        elements = []

        for worksheet in workbook.worksheets:
            cells = []
            values_sheet = values_workbook[worksheet.title]

            # A sheet that has saved results for at least one formula was calculated
            # before saving, so a formula with no saved value produced an empty result
            # (for example IF(...,"",...)) and is skipped instead of indexed as raw text.
            sheet_calculated = any(
                cell.data_type == "f" and values_sheet[cell.coordinate].value is not None
                for row in worksheet.iter_rows()
                for cell in row
            )

            for row in worksheet.iter_rows():
                for cell in row:
                    if cell.value is None:
                        continue

                    value = cell.value
                    entry_extra = {}
                    if cell.data_type == "f":
                        cached = values_sheet[cell.coordinate].value
                        entry_extra["formula"] = str(cell.value)
                        if cached is not None:
                            value = cached
                        elif sheet_calculated:
                            continue

                    if isinstance(value, datetime) and value.time() == datetime.min.time():
                        value = value.date().isoformat()

                    cells.append(
                        {
                            "coordinate": cell.coordinate,
                            "row": cell.row,
                            "column": cell.column,
                            "value": value,
                            "data_type": cell.data_type,
                            **entry_extra,
                        }
                    )

            merged_ranges = [
                str(cell_range)
                for cell_range in worksheet.merged_cells.ranges
            ]

            elements.append(
                ExtractedElement(
                    element_type="worksheet",
                    text=worksheet.title,
                    metadata={
                        "sheet": worksheet.title,
                        "max_row": worksheet.max_row,
                        "max_column": worksheet.max_column,
                        "merged_ranges": merged_ranges,
                    },
                    content={
                        "cells": cells,
                    },
                )
            )

        metadata = {
            "sheet_count": len(workbook.worksheets),
            "sheet_names": workbook.sheetnames,
        }

        workbook.close()
        values_workbook.close()

        return ExtractedDocument(
            source_type="xlsx",
            elements=elements,
            metadata=metadata,
        )
