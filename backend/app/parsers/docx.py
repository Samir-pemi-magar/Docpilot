import re
from typing import BinaryIO

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

from app.parsers.base import DocumentParser
from app.parsers.models import ExtractedDocument, ExtractedElement

_TABLE_TITLE = re.compile(r"^Table\s+\d+\b", re.IGNORECASE)


def _clean(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


class DOCXParser(DocumentParser):
    """Parser for Microsoft Word documents.

    Walks the document body in order so paragraphs, headings and tables stay
    in their original sequence. Tables are emitted in the same shape the PDF
    parser uses (``element_type="table"``, rows joined with " | ", and
    ``content={"rows": ...}``) so the structured chunker handles them.
    """

    def parse(self, file: BinaryIO) -> ExtractedDocument:
        document = Document(file)
        elements: list[ExtractedElement] = []
        paragraph_index = 0
        table_index = 0
        last_heading: str | None = None

        for child in document.element.body.iterchildren():
            tag = child.tag.rsplit("}", 1)[-1]

            if tag == "p":
                paragraph = Paragraph(child, document)
                paragraph_index += 1
                text = paragraph.text.strip()
                if not text:
                    continue
                style_name = paragraph.style.name if paragraph.style is not None else ""
                if style_name.startswith("Heading"):
                    element_type = "heading"
                    last_heading = text
                else:
                    element_type = "paragraph"
                elements.append(
                    ExtractedElement(
                        element_type=element_type,
                        text=text,
                        metadata={"paragraph_index": paragraph_index, "style": style_name},
                    )
                )

            elif tag == "tbl":
                table = Table(child, document)
                table_index += 1
                rows = [[_clean(cell.text) for cell in row.cells] for row in table.rows]
                rendered = [" | ".join(r) for r in rows if any(r)]
                if not rendered:
                    continue
                first_header = rows[0][0] if rows and rows[0] else ""
                table_title = first_header if _TABLE_TITLE.match(first_header) else last_heading
                elements.append(
                    ExtractedElement(
                        element_type="table",
                        text="\n".join(rendered),
                        metadata={
                            "table_index": table_index,
                            "table_title": table_title,
                            "section_heading": last_heading,
                        },
                        content={"rows": rows},
                    )
                )

        metadata = {
            "paragraph_count": len(document.paragraphs),
            "table_count": len(document.tables),
        }
        return ExtractedDocument(source_type="docx", elements=elements, metadata=metadata)
